"""Runs the full daily risk report.

    python main.py            # real data from Yahoo Finance (cached after first run)
    python main.py --refresh  # re-download the data
    python main.py --demo     # synthetic data, no internet needed
"""
import argparse
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

import config as cfg
import risk_engine as re_
from data import load_prices

OUT = os.path.join(os.path.dirname(__file__), "output")


def build_report(demo=False, refresh=False):
    prices = load_prices(demo=demo, refresh=refresh)
    returns_all = prices.pct_change().dropna()
    returns = returns_all.tail(cfg.LOOKBACK)
    pf = re_.Portfolio(prices)
    scale = np.sqrt(cfg.HORIZON_DAYS)

    methods = {
        "Historical simulation": re_.historical_var(pf, returns),
        "Parametric (delta-normal)": re_.parametric_var(pf, returns),
        "Monte Carlo (full revaluation)": re_.monte_carlo_var(pf, returns),
    }
    summary = pd.DataFrame(
        [{"Method": m, "VaR 99% 1d": v, "ES 97.5% 1d": e,
          f"VaR 99% {cfg.HORIZON_DAYS}d": v * scale, f"ES 97.5% {cfg.HORIZON_DAYS}d": e * scale}
         for m, (v, e) in methods.items()])

    greeks = pf.greeks()
    positions = pd.concat([
        pd.DataFrame({"Position": pf.equity.index, "Asset class": "Equity", "Market value (INR)": pf.equity.values}),
        pd.DataFrame({"Position": pf.fx.index, "Asset class": "FX", "Market value (INR)": pf.fx.values}),
        pd.DataFrame({"Position": greeks["Option"], "Asset class": "Options", "Qty": greeks["Qty"],
                      "Strike": greeks["Strike"], "Days to expiry": [o["days"] for o in pf.options],
                      "Market value (INR)": greeks["Value"]}),
    ], ignore_index=True)
    positions = positions.astype({"Qty": "Int64", "Days to expiry": "Int64"})[
        ["Position", "Asset class", "Qty", "Strike", "Days to expiry", "Market value (INR)"]]
    bt = re_.backtest(pf, returns_all)
    breaches = int(bt["Breach"].sum())
    lr, p_value = re_.kupiec_test(len(bt), breaches)
    bt_summary = pd.DataFrame({
        "Metric": ["Days tested", "Breaches", "Expected breaches", "Kupiec LR", "Kupiec p-value",
                   "Model rejected at 5%?", "Basel traffic light"],
        "Value": [len(bt), breaches, round(len(bt) * (1 - cfg.VAR_CONF), 1), round(lr, 3), round(p_value, 4),
                  "Yes" if p_value < 0.05 else "No", re_.traffic_light(breaches)]})

    sheets = {
        "Summary": summary,
        "Positions": positions,
        "Greeks": pf.greeks(),
        "Component VaR": re_.component_var(pf, returns).reset_index(names="Risk factor"),
        "Limits": re_.limit_report(pf, returns),
        "Backtest summary": bt_summary,
        "Backtest daily": bt.reset_index(names="Date"),
        "Stress tests": re_.stress_tests(pf, prices).reset_index(names="Scenario"),
        "PnL attribution": re_.pnl_attribution(pf, prices).reset_index(names="Date"),
    }

    os.makedirs(OUT, exist_ok=True)
    write_excel(sheets, os.path.join(OUT, "risk_report.xlsx"))
    plot_distribution(pf.pnl(returns), methods["Historical simulation"], os.path.join(OUT, "pnl_distribution.png"))
    plot_backtest(bt, os.path.join(OUT, "backtest.png"))

    pd.options.display.float_format = "{:,.2f}".format
    pd.options.display.width = 200
    print(f"\nAs of {prices.index[-1].date()} | NIFTY {pf.spot:,.0f} | implied vol {pf.vol:.1%}"
          + (" | SYNTHETIC DEMO DATA" if demo else ""))
    for name in ["Summary", "Limits", "Backtest summary", "Stress tests", "PnL attribution"]:
        print(f"\n--- {name} ---\n{sheets[name].to_string(index=False)}")
    print(f"\nReport written to {OUT}")
    return sheets


def write_excel(sheets, path):
    red = PatternFill("solid", fgColor="F8CBAD")
    amber = PatternFill("solid", fgColor="FFE699")
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, df in sheets.items():
            df.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value)) for c in col) + 3
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0%" if "%" in str(ws.cell(1, cell.column).value) or \
                            ws.cell(1, cell.column).value == "Utilisation" else "#,##0.00"
                    if cell.value in ("BREACH", True):
                        cell.fill = red
                    elif cell.value == "WARNING":
                        cell.fill = amber


def plot_distribution(pnl, hist, path):
    var, es = hist
    fig, ax = plt.subplots(figsize=(9, 4.5))
    ax.hist(pnl, bins=60, color="#4C72B0", alpha=0.85)
    ax.axvline(-var, color="#C44E52", linestyle="--", label=f"99% VaR: {var:,.0f}")
    ax.axvline(-es, color="#8C1C13", linestyle="-", label=f"97.5% ES: {es:,.0f}")
    ax.set(title="1-day P&L distribution (historical simulation)", xlabel="P&L (INR)", ylabel="Days")
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def plot_backtest(bt, path):
    fig, ax = plt.subplots(figsize=(9, 4.5))
    ax.bar(bt.index, bt["P&L"], color="#4C72B0", width=1.0, label="Daily P&L")
    ax.plot(bt.index, -bt["VaR 99%"], color="#C44E52", label="-VaR 99%")
    hits = bt[bt["Breach"]]
    ax.scatter(hits.index, hits["P&L"], color="black", zorder=3, label=f"Breaches ({len(hits)})")
    ax.set(title="VaR backtest", ylabel="INR")
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--demo", action="store_true", help="use synthetic data")
    ap.add_argument("--refresh", action="store_true", help="re-download prices")
    args = ap.parse_args()
    build_report(demo=args.demo, refresh=args.refresh)
